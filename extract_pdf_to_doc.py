import fitz  # PyMuPDF
import sys
import os
from docx import Document

def extract_pdf_content(pdf_path):
    try:
        # Open the PDF
        doc = fitz.open(pdf_path)
        print(f"Opened PDF: {pdf_path}")

        # Prepare output .docx file
        base_name = os.path.splitext(os.path.basename(pdf_path))[0]
        output_file = f"{base_name}_extracted.docx"

        document = Document()

        # Header section
        document.add_heading("PDF Extraction Output", level=1)
        document.add_paragraph(f"Source File: {pdf_path}")
        document.add_paragraph(f"Total Pages: {len(doc)}")
        document.add_paragraph("-" * 50)

        # Iterate through PDF pages
        for page_num in range(len(doc)):
            page = doc[page_num]

            document.add_heading(f"Page {page_num + 1}", level=2)

            # Extract text
            text = page.get_text()
            document.add_heading("Page Text", level=3)
            document.add_paragraph(text if text.strip() else "(No extractable text)")

            # Extract annotations
            document.add_heading("Annotations", level=3)

            annot = page.first_annot
            if not annot:
                document.add_paragraph("No annotations found.")
            else:
                i = 1
                while annot:
                    subtype = annot.type[0]
                    content = annot.info.get("content", "")
                    document.add_paragraph(f"Annotation {i}: ({subtype}) {content}")
                    annot = annot.next
                    i += 1

            document.add_paragraph("-" * 50)

        # Save the Word file
        document.save(output_file)
        doc.close()

        print(f"\nExtraction completed. Output saved to: {output_file}")

    except Exception as e:
        print(f"Error: {e}")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python extract_pdf_to_doc.py <path_to_pdf>")
    else:
        extract_pdf_content(sys.argv[1])
