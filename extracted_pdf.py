import fitz  # PyMuPDF
import pytesseract
from PIL import Image
import sys
import os
import io
import re
import json
from docx import Document


# ----------------- CLEAN TEXT FOR DOCX -----------------
def clean_text(text):
    if not text:
        return ""
    return re.sub(r'[^\x09\x0A\x0D\x20-\uD7FF\uE000-\uFFFD]', '', text)


# ----------------- OCR IMAGE -----------------
def ocr_image(pix):
    try:
        img_bytes = pix.tobytes("png")
        img = Image.open(io.BytesIO(img_bytes))
        return clean_text(pytesseract.image_to_string(img))
    except Exception as e:
        return f"(OCR failed: {e})"


# ----------------- VECTOR GRAPHICS EXTRACTION -----------------
def extract_vector_graphics(page):
    drawings = []
    for item in page.get_drawings():
        entry = {
            "type": item.get("type", ""),
            "rect": list(item.get("rect", [])),
            "color": item.get("color", None),
            "fill": item.get("fill", None),
            "linewidth": item.get("width", None),
            "points": item.get("points", None),
        }
        drawings.append(entry)
    return drawings


# ----------------- DETECT IF COLOR IS RED -----------------
def is_red_color(color_tuple):
    """PyMuPDF gives colors as floats (0–1). Detect strong red."""
    if not color_tuple:
        return False
    r, g, b = color_tuple[:3]
    return r > 0.8 and g < 0.3 and b < 0.3


# ----------------- RED MARKUP ANNOTATION EXTRACTION -----------------
def extract_red_annotations(page):
    results = []
    annot = page.first_annot

    while annot:
        subtype = annot.type[0] if annot.type else "Unknown"
        info = annot.info or {}

        content = clean_text(info.get("content", ""))
        title = clean_text(info.get("title", ""))

        border_color = annot.colors.get("stroke") if annot.colors else None
        fill_color = annot.colors.get("fill") if annot.colors else None

        if is_red_color(border_color) or is_red_color(fill_color):
            results.append({
                "subtype": subtype,
                "author": title,
                "content": content,
                "border_color": border_color,
                "fill_color": fill_color,
                "rect": list(annot.rect),
            })

        annot = annot.next

    return results


# ----------------- MAIN EXTRACTION PIPELINE -----------------
def extract_pdf_content(pdf_path):
    try:
        doc = fitz.open(pdf_path)
        print(f"Opened PDF: {pdf_path}")

        base_name = os.path.splitext(os.path.basename(pdf_path))[0]
        output_docx = f"{base_name}_enhanced_extraction.docx"
        output_vectors = f"{base_name}_vectors.json"
        output_red_annots = f"{base_name}_red_annotations.json"

        document = Document()
        document.add_heading("Enhanced PDF Extraction Output", level=1)
        document.add_paragraph(f"Source File: {pdf_path}")
        document.add_paragraph(f"Total Pages: {len(doc)}")
        document.add_paragraph("-" * 60)

        all_drawings = {}
        all_red_annotations = {}

        # -------------- PROCESS PAGES --------------
        for page_num, page in enumerate(doc):
            document.add_heading(f"Page {page_num + 1}", level=2)

            # TEXT EXTRACTION
            text = page.get_text()

            if not text.strip():
                document.add_paragraph("No vector text — using OCR")
                pix = page.get_pixmap(dpi=300)
                text = ocr_image(pix)

            text = clean_text(text)
            document.add_heading("Page Text", level=3)
            document.add_paragraph(text)

            # VECTOR GRAPHICS
            drawings = extract_vector_graphics(page)
            all_drawings[f"page_{page_num+1}"] = drawings

            document.add_heading("Detected Lines / Shapes / Symbols", level=3)
            if drawings:
                document.add_paragraph(f"{len(drawings)} vectors detected.")
            else:
                document.add_paragraph("No vector graphics found.")

            # RED MARKUP EXTRACTION
            red_annots = extract_red_annotations(page)
            all_red_annotations[f"page_{page_num+1}"] = red_annots

            document.add_heading("Red Markup Comments", level=3)
            if red_annots:
                for idx, a in enumerate(red_annots, 1):
                    document.add_paragraph(
                        f"{idx}. [{a['subtype']}] {a['content']}\n"
                        f"Author: {a['author']}\n"
                        f"Color: {a['border_color']}\n"
                        f"Rect: {a['rect']}"
                    )
            else:
                document.add_paragraph("No red annotations detected.")

        # ----------- SAVE OUTPUT FILES -----------
        document.save(output_docx)

        with open(output_vectors, "w", encoding="utf-8") as f:
            json.dump(all_drawings, f, indent=4)

        with open(output_red_annots, "w", encoding="utf-8") as f:
            json.dump(all_red_annotations, f, indent=4)

        print(f"Extraction completed.\nDOCX saved as: {output_docx}\n"
              f"Vector graphics JSON: {output_vectors}\n"
              f"Red annotations JSON: {output_red_annots}")

        return {
            "docx": output_docx,
            "vectors_json": output_vectors,
            "red_annots_json": output_red_annots
        }

    except Exception as e:
        print(f"Error processing PDF: {e}")
        return None


# ----------------- COMMAND-LINE INTERFACE -----------------
if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python extract_pdf.py <PDF_FILE>")
    else:
        extract_pdf_content(sys.argv[1])
